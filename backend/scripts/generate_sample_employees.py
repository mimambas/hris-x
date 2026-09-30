"""Generator data contoh 500 karyawan untuk demo impor Excel (Sprint 3, CHR-007).

Jalankan dari direktori backend/:
    ../.venv/bin/python scripts/generate_sample_employees.py

Menghasilkan backend/sample_data/karyawan_500.xlsx dengan kolom sesuai
template impor (lihat app/services/imports.py :: HEADER_ROW). NIK 16 digit
unik format wilayah+DDMMYY+serial; nama Indonesia realistis.

Opsi:
    --count N     jumlah baris (default 500)
    --output PATH berkas keluaran
    --seed N      seed RNG untuk hasil reproduksibel
"""

from __future__ import annotations

import argparse
import random
import sys
from datetime import date
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services.imports import HEADER_ROW  # noqa: E402

NAMA_DEPAN_PRIA = [
    "Agus", "Andi", "Budi", "Dedi", "Eko", "Fajar", "Hendra", "Irfan", "Joko",
    "Rudi", "Yudi", "Danu", "Bima", "Galih", "Rizky", "Fikri", "Ilham",
    "Bagus", "Dimas", "Yoga", "Aditya", "Putra", "Wahyu", "Teguh", "Hadi",
    "Slamet", "Sutrisno", "Haryanto", "Nugroho", "Pratama", "Setiawan",
]
NAMA_DEPAN_WANITA = [
    "Sari", "Dewi", "Rina", "Maya", "Putri", "Ayu", "Nita", "Lina", "Fitri",
    "Wulan", "Ratna", "Indah", "Novita", "Dian", "Rani", "Sinta", "Eka",
    "Yuni", "Sri", "Tania", "Melati", "Kirana", "Sekar", " Laras",
]
NAMA_BELAKANG = [
    "Santoso", "Wijaya", "Pratama", "Lestari", "Kusuma", "Saputra", "Nugroho",
    "Hidayat", "Rahmawati", "Setiawan", "Hartono", "Gunawan", "Siregar",
    "Nasution", "Puspita", "Anggraini", "Maharani", "Firmansyah", "Ramadhan",
    "Saputra", "Wibowo", "Kurniawan", "Handoko", "Sutanto", "Halim",
    "Sanjaya", "Pangestu", "Wicaksono", "Utama", "Permana", "Hakim",
]
WILAYAH = ["317401", "317402", "317403", "327301", "327302", "337401",
           "357801", "327101", "317404", "327303"]
BANKS = ["BCA", "BRI", "Mandiri", "BNI", "BTN"]
PTKP_POOL = (["TK/0"] * 40 + ["K/0"] * 20 + ["TK/1"] * 10 + ["K/1"] * 10
             + ["TK/2"] * 5 + ["K/2"] * 5 + ["KI/0"] * 5 + ["TK/3"] * 5)

LEGAL_ENTITIES = ["PT Hashiru Teknologi", "PT Hashiru Distribusi"]
ORG_UNITS = ["Tim Backend", "Departemen Engineering", "Departemen Logistik"]
LOKASI = ["Kantor Pusat Jakarta", "Gudang Bekasi"]
JOBS = [
    # (kode, bobot, gaji_min, gaji_max)
    ("STF", 60, 5_000_000, 9_000_000),
    ("SPV", 25, 10_000_000, 14_000_000),
    ("MGR", 15, 16_000_000, 25_000_000),
]
POSISI = {
    "STF": ["Staff Administrasi", "Staff Operasional", "Backend Engineer",
            "Staff Gudang", "Customer Service"],
    "SPV": ["Supervisor Operasional", "Supervisor Engineering",
            "Supervisor Gudang"],
    "MGR": ["Engineering Manager", "Operations Manager", "Finance Manager"],
}


def _pick_job(rng):
    total = sum(b for _, b, _, _ in JOBS)
    x = rng.uniform(0, total)
    acc = 0
    for kode, bobot, gmin, gmax in JOBS:
        acc += bobot
        if x <= acc:
            gaji = rng.randrange(gmin, gmax + 1, 100_000)
            tunj = rng.choice([500_000, 1_000_000, 1_500_000, 2_000_000,
                               2_500_000, 3_000_000])
            return kode, gaji, tunj
    kode, _, gmin, gmax = JOBS[0]
    return kode, gmin, 500_000


def generate(count: int, seed: int) -> list[dict]:
    rng = random.Random(seed)
    used_nik, used_email = set(), set()
    rows = []
    for i in range(count):
        pria = rng.random() < 0.55
        depan = rng.choice(NAMA_DEPAN_PRIA if pria else NAMA_DEPAN_WANITA).strip()
        belakang = rng.choice(NAMA_BELAKANG)
        nama = f"{depan} {belakang}"

        # NIK: wilayah(6) + DDMMYY lahir + serial(4), unik.
        lahir = date(rng.randint(1970, 2004), rng.randint(1, 12), rng.randint(1, 28))
        while True:
            nik = (rng.choice(WILAYAH) + lahir.strftime("%d%m%y")
                   + f"{rng.randint(0, 9999):04d}")
            if nik not in used_nik:
                used_nik.add(nik)
                break

        base_email = f"{depan}.{belakang}".lower().replace(" ", "")
        n = 0
        while True:
            suffix = "" if n == 0 else str(n)
            email = f"{base_email}{suffix}@contoh.id"
            n += 1
            if email not in used_email:
                used_email.add(email)
                break

        job_code, gaji, tunj = _pick_job(rng)
        masuk = date(rng.randint(2020, 2026), rng.randint(1, 12), rng.randint(1, 28))
        if masuk > date(2026, 9, 30):
            masuk = date(2026, 9, 30)
        is_pkwt = rng.random() < 0.45
        if is_pkwt:
            dur = rng.choice([12, 12, 24])
            y = masuk.year + (masuk.month - 1 + dur) // 12
            m = (masuk.month - 1 + dur) % 12 + 1
            contract_end = date(y, m, min(masuk.day, 28))
        else:
            contract_end = None

        rows.append({
            "nik": nik,
            "nama": nama,
            "email": email,
            "tgl_lahir": lahir,
            "npwp": (f"{rng.randint(10**15, 10**16 - 1)}"
                     if rng.random() < 0.85 else ""),
            "ptkp": rng.choice(PTKP_POOL),
            "bpjs_kes": (f"{rng.randint(10**12, 10**13 - 1)}"
                         if rng.random() < 0.9 else ""),
            "bpjs_tk": (f"{rng.randint(10**10, 10**11 - 1)}"
                        if rng.random() < 0.9 else ""),
            "bank": rng.choice(BANKS),
            "no_rekening": f"{rng.randint(10**9, 10**10 - 1)}",
            "legal_entity": rng.choice(LEGAL_ENTITIES),
            "org_unit": rng.choice(ORG_UNITS),
            "job_code": job_code,
            "position": rng.choice(POSISI[job_code]),
            "tgl_masuk": masuk,
            "contract_type": "PKWT" if is_pkwt else "PKWTT",
            "contract_end": contract_end or "",
            "gaji_pokok": gaji,
            "tunjangan_tetap": tunj,
            "lokasi": rng.choice(LOKASI),
        })
    return rows


def write_xlsx(rows: list[dict], output: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Karyawan"
    for ci, key in enumerate(HEADER_ROW, start=1):
        ws.cell(row=1, column=ci, value=key)
    date_fmt = "DD/MM/YYYY"
    for ri, row in enumerate(rows, start=2):
        for ci, key in enumerate(HEADER_ROW, start=1):
            val = row[key]
            cell = ws.cell(row=ri, column=ci, value=val if val != "" else None)
            if isinstance(val, date):
                cell.number_format = date_fmt
            if key in ("nik", "npwp", "bpjs_kes", "bpjs_tk", "no_rekening"):
                cell.number_format = "@"  # teks: cegah notasi ilmiah
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--count", type=int, default=500)
    ap.add_argument("--output", default="sample_data/karyawan_500.xlsx")
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()
    rows = generate(args.count, args.seed)
    out = Path(args.output)
    if not out.is_absolute():
        out = Path(__file__).resolve().parent.parent / out
    write_xlsx(rows, out)
    print(f"Ditulis {len(rows)} baris -> {out}")


if __name__ == "__main__":
    main()

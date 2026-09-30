"""Impor karyawan dari Excel: template, dry-run, commit atomik (CHR-007).

Alur:
- GET template  -> build_template()
- POST dry-run -> validate_all() -> laporan per baris, TANPA tulis DB
- POST commit  -> validate_all() dulu; bila ada 1 baris invalid -> 422 +
  laporan lengkap, DB tak tersentuh; bila semua valid -> tulis semua dalam
  satu transaksi (atomic) + audit channel="import".

Kolom dicocokkan: legal_entity/org_unit/lokasi by NAMA (case-insensitive),
job by kode. NIK 16 digit & email unik per tenant + unik di dalam file.
"""

from __future__ import annotations

import re
from datetime import date
from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (
    CompInfo,
    Contract,
    ContractInfo,
    Employment,
    Job,
    JobInfo,
    LegalEntity,
    LegalEntityInfo,
    Location,
    LocationInfo,
    OrgUnit,
    OrgUnitInfo,
    Person,
    Position,
)
from app.schemas.schemas import (
    ImportCommitResponse,
    ImportDryRunResponse,
    ImportRowError,
    ImportRowReport,
)
from app.services import effective_dating as ed
from app.services.audit import write_audit
from app.services.contracts import ensure_number_unique, get_policy, months_between
from app.services.validation import (
    parse_id_date,
    validate_bank_account,
    validate_birth_date,
    validate_bpjs,
    validate_email,
    validate_nik,
    validate_npwp,
    validate_ptkp,
)

# (key, wajib, panduan)
COLUMNS: list[tuple[str, bool, str]] = [
    ("nik", True, "16 digit angka, unik per tenant"),
    ("nama", True, "Nama lengkap karyawan"),
    ("email", False, "Opsional; format valid & unik per tenant"),
    ("tgl_lahir", False, "Format DD/MM/YYYY; tidak boleh masa depan"),
    ("npwp", False, "16 digit (format baru); titik/strip/spasi diabaikan"),
    ("ptkp", False, "TK/0, TK/1, TK/2, TK/3, K/0..K/3, KI/0..KI/3. Default TK/0"),
    ("bpjs_kes", False, "Nomor BPJS Kesehatan, 10-16 digit"),
    ("bpjs_tk", False, "Nomor BPJS Ketenagakerjaan, 10-16 digit"),
    ("bank", False, "Nama bank, mis. BCA / BRI / Mandiri / BNI"),
    ("no_rekening", False, "Nomor rekening, 4-32 digit"),
    ("legal_entity", True, "NAMA legal entity persis seperti di sistem"),
    ("org_unit", True, "NAMA unit organisasi persis seperti di sistem"),
    ("job_code", True, "Kode jabatan, mis. STF / SPV / MGR"),
    ("position", True, "Nama posisi; dibuat otomatis bila belum ada"),
    ("tgl_masuk", True, "Tanggal masuk kerja, format DD/MM/YYYY"),
    ("contract_type", True, "PKWT atau PKWTT"),
    ("contract_end", False, "WAJIB bila PKWT; format DD/MM/YYYY"),
    ("gaji_pokok", True, "Integer rupiah, mis. 8000000"),
    ("tunjangan_tetap", True, "Integer rupiah, mis. 2000000"),
    ("lokasi", False, "NAMA lokasi persis; kosong = lokasi pertama (abjad)"),
]

HEADER_ROW = [key for key, _, _ in COLUMNS]
REQUIRED_KEYS = [key for key, req, _ in COLUMNS if req]


class ImportValidationError(Exception):
    """Commit dibatalkan karena ada baris invalid; membawa laporan dry-run."""

    def __init__(self, report: ImportDryRunResponse):
        self.report = report
        super().__init__(
            f"Impor dibatalkan: {report.invalid_rows} dari {report.total_rows} "
            f"baris tidak valid. DB tidak berubah."
        )


def _as_text(value) -> str:
    """Normalisasi sel Excel ke teks (angka 16 digit dari Excel bisa float)."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_money(value) -> int:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError("Nominal wajib diisi")
    if isinstance(value, bool):
        raise ValueError("Nominal tidak valid")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            raise ValueError("Nominal harus bilangan bulat rupiah")
        amount = int(value)
    else:
        text = re.sub(r"[.\s]", "", str(value).strip())
        if not re.fullmatch(r"\d+", text or ""):
            raise ValueError("Nominal harus angka bulat rupiah")
        amount = int(text)
    if amount < 0:
        raise ValueError("Nominal tidak boleh negatif")
    return amount


# ------------------------------------------------------------------ template
def build_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Karyawan"
    fill = PatternFill("solid", fgColor="1F4E5F")
    font = Font(bold=True, color="FFFFFF")
    for ci, key in enumerate(HEADER_ROW, start=1):
        cell = ws.cell(row=1, column=ci, value=key)
        cell.font = font
        cell.fill = fill
    examples = [
        ["3174010101900001", "Budi Santoso", "budi.santoso@contoh.id",
         "01/01/1990", "1234567890123456", "TK/0", "0001234567890",
         "12345678901", "BCA", "1234567890", "PT Hashiru Teknologi",
         "Tim Backend", "STF", "Backend Engineer", "01/03/2024", "PKWTT",
         "", 8000000, 2000000, "Kantor Pusat Jakarta"],
        ["3174010202920002", "Sari Wijaya", "sari.wijaya@contoh.id",
         "02/02/1992", "", "K/0", "", "", "Mandiri", "0987654321",
         "PT Hashiru Teknologi", "Tim Backend", "STF", "Backend Engineer",
         "10/01/2026", "PKWT", "09/01/2027", 5000000, 500000,
         "Kantor Pusat Jakarta"],
    ]
    for ri, row in enumerate(examples, start=2):
        for ci, val in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=val)
    widths = {"nik": 20, "nama": 22, "email": 28, "tgl_lahir": 13, "npwp": 20,
              "ptkp": 8, "bpjs_kes": 18, "bpjs_tk": 18, "bank": 12,
              "no_rekening": 16, "legal_entity": 24, "org_unit": 24,
              "job_code": 11, "position": 22, "tgl_masuk": 13,
              "contract_type": 14, "contract_end": 15, "gaji_pokok": 13,
              "tunjangan_tetap": 16, "lokasi": 22}
    for ci, key in enumerate(HEADER_ROW, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = (
            widths.get(key, 15)
        )

    guide = wb.create_sheet("Panduan")
    guide.append(["Kolom", "Wajib", "Keterangan"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for key, req, desc in COLUMNS:
        guide.append([key, "Ya" if req else "Tidak", desc])
    guide.append([])
    guide.append(["Catatan"])
    for note in [
        "1. Hapus 2 baris contoh di sheet Karyawan sebelum mengimpor data asli.",
        "2. Tanggal: format DD/MM/YYYY (atau biarkan format Tanggal Excel).",
        "3. legal_entity / org_unit / lokasi diisi NAMA persis seperti di sistem "
        "(bukan kode); huruf besar-kecil diabaikan.",
        "4. NIK 16 digit & email harus unik (per tenant dan di dalam file).",
        "5. Dry-run memvalidasi SEMUA baris tanpa menulis database; commit "
        "hanya jalan bila 100% baris valid (atomic: gagal satu = gagal semua).",
        "6. Impor memakai event katalog: hire / Impor data massal.",
    ]:
        guide.append([note])
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 8
    guide.column_dimensions["C"].width = 70

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ parsing
def parse_workbook(data: bytes) -> list[dict]:
    """Urai workbook -> [{row_number, values}]. ValueError bila file fatal."""
    try:
        wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise ValueError(f"Berkas bukan file Excel yang valid: {e}")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Berkas Excel kosong")
    headers = [_as_text(c).lower() for c in rows[0]]
    missing = [k for k in REQUIRED_KEYS if k not in headers]
    if missing:
        raise ValueError(f"Kolom wajib hilang: {', '.join(missing)}")
    col_idx = {k: headers.index(k) for k in HEADER_ROW if k in headers}
    out = []
    for idx, r in enumerate(rows[1:], start=2):
        r = r or []
        vals = {k: (r[i] if i < len(r) else None) for k, i in col_idx.items()}
        if all(v is None or (isinstance(v, str) and not v.strip())
               for v in vals.values()):
            continue  # lewati baris kosong
        out.append({"row_number": idx, "values": vals})
    if not out:
        raise ValueError("Tidak ada baris data di berkas Excel")
    return out


# ------------------------------------------------------------------ validasi
class _Ctx:
    def __init__(self, db: Session, tenant_id):
        self.db = db
        self.tenant_id = tenant_id
        self.legal_entities = _name_map(db, tenant_id, LegalEntityInfo,
                                        "legal_entity_id")
        self.org_units = _name_map(db, tenant_id, OrgUnitInfo, "org_unit_id")
        self.locations = _name_map(db, tenant_id, LocationInfo, "location_id")
        self.jobs = {
            (j.code or "").strip().lower(): j.id
            for j in db.execute(
                select(Job).where(Job.tenant_id == tenant_id)
            ).scalars().all()
        }
        self.existing_niks = {
            n for (n,) in db.execute(
                select(Person.nik).where(Person.tenant_id == tenant_id)
            ).all()
        }
        self.existing_emails = {
            e.lower() for (e,) in db.execute(
                select(Person.email).where(
                    Person.tenant_id == tenant_id, Person.email.is_not(None)
                )
            ).all()
        }
        names = sorted(self.locations.keys())
        self.default_location_id = self.locations[names[0]] if names else None
        self.policy = get_policy(db, tenant_id)


def _name_map(db: Session, tenant_id, info_model, identity_field: str) -> dict:
    """nama (lower) -> identity id, dari versi info terbaru tiap identitas."""
    rows = (
        db.execute(
            select(info_model)
            .where(info_model.tenant_id == tenant_id)
            .order_by(info_model.valid_from.desc(), info_model.seq_no.desc())
        )
        .scalars()
        .all()
    )
    out, seen = {}, set()
    for r in rows:
        ident = getattr(r, identity_field)
        if ident in seen:
            continue
        seen.add(ident)
        out[(getattr(r, "name") or "").strip().lower()] = ident
    return out


def _validate_row(ctx: _Ctx, row_number: int, values: dict,
                  seen_niks: dict, seen_emails: dict
                  ) -> tuple[list[ImportRowError], dict]:
    errors: list[ImportRowError] = []
    clean: dict = {}

    def err(field: str, message: str) -> None:
        errors.append(ImportRowError(field=field, message=message))

    # --- NIK ---
    try:
        nik = validate_nik(_as_text(values.get("nik")))
    except ValueError as e:
        err("nik", str(e))
        nik = None
    if nik:
        if nik in ctx.existing_niks:
            err("nik", "NIK sudah terdaftar di tenant ini")
        elif nik in seen_niks:
            err("nik", f"NIK duplikat dengan baris {seen_niks[nik]} di file ini")
        else:
            seen_niks[nik] = row_number
    clean["nik"] = nik

    # --- nama ---
    nama = _as_text(values.get("nama"))
    if not nama:
        err("nama", "Nama wajib diisi")
    clean["nama"] = nama

    # --- email ---
    try:
        email = validate_email(_as_text(values.get("email")))
    except ValueError as e:
        err("email", str(e))
        email = None
    if email:
        key = email.lower()
        if key in ctx.existing_emails:
            err("email", "Email sudah terdaftar di tenant ini")
        elif key in seen_emails:
            err("email", f"Email duplikat dengan baris {seen_emails[key]} di file ini")
        else:
            seen_emails[key] = row_number
    clean["email"] = email

    # --- identitas opsional ---
    try:
        clean["tgl_lahir"] = validate_birth_date(parse_id_date(values.get("tgl_lahir")))
    except ValueError as e:
        err("tgl_lahir", str(e))
        clean["tgl_lahir"] = None
    try:
        clean["npwp"] = validate_npwp(_as_text(values.get("npwp")))
    except ValueError as e:
        err("npwp", str(e))
        clean["npwp"] = None
    try:
        clean["ptkp"] = validate_ptkp(_as_text(values.get("ptkp")))
    except ValueError as e:
        err("ptkp", str(e))
        clean["ptkp"] = "TK/0"
    try:
        clean["bpjs_kes"] = validate_bpjs(_as_text(values.get("bpjs_kes")),
                                         "BPJS Kesehatan")
    except ValueError as e:
        err("bpjs_kes", str(e))
        clean["bpjs_kes"] = None
    try:
        clean["bpjs_tk"] = validate_bpjs(_as_text(values.get("bpjs_tk")),
                                        "BPJS Ketenagakerjaan")
    except ValueError as e:
        err("bpjs_tk", str(e))
        clean["bpjs_tk"] = None
    clean["bank"] = _as_text(values.get("bank")) or None
    try:
        clean["no_rekening"] = validate_bank_account(_as_text(values.get("no_rekening")))
    except ValueError as e:
        err("no_rekening", str(e))
        clean["no_rekening"] = None

    # --- referensi master ---
    le_name = _as_text(values.get("legal_entity"))
    le_id = ctx.legal_entities.get(le_name.lower())
    if not le_name:
        err("legal_entity", "Legal entity wajib diisi")
    elif le_id is None:
        err("legal_entity", f"Legal entity '{le_name}' tidak ditemukan di sistem")
    clean["legal_entity_id"] = le_id

    ou_name = _as_text(values.get("org_unit"))
    ou_id = ctx.org_units.get(ou_name.lower())
    if not ou_name:
        err("org_unit", "Unit organisasi wajib diisi")
    elif ou_id is None:
        err("org_unit", f"Unit organisasi '{ou_name}' tidak ditemukan di sistem")
    clean["org_unit_id"] = ou_id

    job_code = _as_text(values.get("job_code"))
    job_id = ctx.jobs.get(job_code.lower())
    if not job_code:
        err("job_code", "Kode jabatan wajib diisi")
    elif job_id is None:
        err("job_code", f"Kode jabatan '{job_code}' tidak ditemukan di sistem")
    clean["job_id"] = job_id

    position = _as_text(values.get("position"))
    if not position:
        err("position", "Posisi wajib diisi")
    clean["position"] = position

    loc_name = _as_text(values.get("lokasi"))
    if loc_name:
        loc_id = ctx.locations.get(loc_name.lower())
        if loc_id is None:
            err("lokasi", f"Lokasi '{loc_name}' tidak ditemukan di sistem")
    else:
        loc_id = ctx.default_location_id
        if loc_id is None:
            err("lokasi", "Tenant belum memiliki lokasi")
    clean["location_id"] = loc_id

    # --- tanggal masuk ---
    try:
        tgl_masuk = parse_id_date(values.get("tgl_masuk"))
    except ValueError as e:
        err("tgl_masuk", str(e))
        tgl_masuk = None
    if tgl_masuk is None and not any(e.field == "tgl_masuk" for e in errors):
        err("tgl_masuk", "Tanggal masuk wajib diisi")
    clean["tgl_masuk"] = tgl_masuk

    # --- kontrak ---
    ctype = _as_text(values.get("contract_type")).upper()
    if ctype not in ("PKWT", "PKWTT"):
        err("contract_type", "Jenis kontrak harus PKWT atau PKWTT")
        ctype = None
    clean["contract_type"] = ctype
    raw_end = values.get("contract_end")
    contract_end = None
    try:
        contract_end = parse_id_date(raw_end)
    except ValueError as e:
        err("contract_end", str(e))
        contract_end = "ERROR"  # penanda: sudah dilaporkan
    if ctype == "PKWT" and contract_end != "ERROR":
        if contract_end is None:
            err("contract_end", "Kontrak PKWT wajib memiliki tanggal berakhir")
        elif tgl_masuk and contract_end <= tgl_masuk:
            err("contract_end", "Tanggal berakhir harus setelah tanggal masuk")
        elif (tgl_masuk
              and months_between(tgl_masuk, contract_end) > ctx.policy.max_pkwt_months):
            err("contract_end",
                f"Durasi PKWT melebihi batas kebijakan ({ctx.policy.max_pkwt_months} bulan)")
    elif ctype == "PKWTT" and isinstance(contract_end, date):
        err("contract_end", "Kontrak PKWTT tidak memakai tanggal berakhir")
    clean["contract_end"] = contract_end if isinstance(contract_end, date) else None

    # --- gaji ---
    for field in ("gaji_pokok", "tunjangan_tetap"):
        try:
            clean[field] = _parse_money(values.get(field))
        except ValueError as e:
            err(field, str(e))
            clean[field] = None

    return errors, clean


def validate_all(db: Session, tenant_id, filename: str, data: bytes
                 ) -> tuple[list[dict], ImportDryRunResponse]:
    """Validasi seluruh baris. Mengembalikan (baris_bersih_valid, laporan)."""
    ctx = _Ctx(db, tenant_id)
    raw_rows = parse_workbook(data)  # ValueError bila file fatal
    seen_niks: dict[str, int] = {}
    seen_emails: dict[str, int] = {}
    cleaned, invalid = [], []
    for r in raw_rows:
        errors, clean = _validate_row(ctx, r["row_number"], r["values"],
                                      seen_niks, seen_emails)
        if errors:
            invalid.append(ImportRowReport(
                row_number=r["row_number"],
                status="invalid",
                errors=errors,
                preview={"nik": _as_text(r["values"].get("nik")),
                         "nama": _as_text(r["values"].get("nama"))},
            ))
        else:
            cleaned.append({"row_number": r["row_number"], **clean})
    report = ImportDryRunResponse(
        filename=filename,
        total_rows=len(raw_rows),
        valid_rows=len(cleaned),
        invalid_rows=len(invalid),
        rows=invalid,
    )
    return cleaned, report


# ------------------------------------------------------------------ commit
def _write_row(db: Session, tenant_id, user, filename: str, c: dict,
               ip: str | None) -> object:
    audit_reason = f"Impor Excel '{filename}' baris {c['row_number']}"

    def _audit(action, object_type, object_id, new_values):
        write_audit(
            db=db, tenant_id=tenant_id, actor_user_id=user.id, action=action,
            object_type=object_type, object_id=object_id, new_values=new_values,
            reason=audit_reason, channel="import", ip=ip,
        )

    person = Person(
        tenant_id=tenant_id, nik=c["nik"], full_name=c["nama"],
        birth_date=c["tgl_lahir"], email=c["email"], npwp=c["npwp"],
        ptkp=c["ptkp"], bpjs_kes_no=c["bpjs_kes"], bpjs_tk_no=c["bpjs_tk"],
        bank_name=c["bank"], bank_account_no=c["no_rekening"],
    )
    db.add(person)
    db.flush()
    _audit("create", "person", person.id,
           {"nik": person.nik, "full_name": person.full_name})

    emp = Employment(
        tenant_id=tenant_id, person_id=person.id,
        legal_entity_id=c["legal_entity_id"], start_date=c["tgl_masuk"],
        end_date=None, status="active",
    )
    db.add(emp)
    db.flush()
    _audit("create", "employment", emp.id,
           {"person_id": str(person.id),
            "start_date": c["tgl_masuk"].isoformat()})

    # Position: pakai yang ada bila (job, unit, nama) cocok; else buat baru.
    pos = (
        db.execute(
            select(Position).where(
                Position.tenant_id == tenant_id,
                Position.job_id == c["job_id"],
                Position.org_unit_id == c["org_unit_id"],
                func.lower(Position.name) == c["position"].lower(),
            )
        )
        .scalars()
        .first()
    )
    if pos is None:
        pos = Position(tenant_id=tenant_id, job_id=c["job_id"],
                       org_unit_id=c["org_unit_id"], name=c["position"])
        db.add(pos)
        db.flush()
        _audit("create", "position", pos.id, {"name": pos.name})

    job_rec = ed.insert_record(
        db=db, tenant_id=tenant_id, model=JobInfo,
        identity_field="employment_id", identity_value=emp.id,
        valid_from=c["tgl_masuk"],
        values={"job_id": c["job_id"], "org_unit_id": c["org_unit_id"],
                "location_id": c["location_id"], "manager_employment_id": None},
        event="hire", event_reason="Impor data massal", created_by=user.id,
        event_applies_to="lifecycle",
    )
    _audit("insert", "job_info", job_rec.id, {"employment_id": str(emp.id)})

    comp_rec = ed.insert_record(
        db=db, tenant_id=tenant_id, model=CompInfo,
        identity_field="employment_id", identity_value=emp.id,
        valid_from=c["tgl_masuk"],
        values={"pay_group": "Bulanan",
                "components": {"gaji_pokok": c["gaji_pokok"],
                               "tunjangan_tetap": c["tunjangan_tetap"]},
                "ptkp": person.ptkp},
        event="hire", event_reason="Penetapan gaji awal", created_by=user.id,
        event_applies_to="lifecycle",
    )
    _audit("insert", "comp_info", comp_rec.id, {"employment_id": str(emp.id)})

    contract = Contract(tenant_id=tenant_id, employment_id=emp.id)
    db.add(contract)
    db.flush()
    number = f"IMP-{c['nik']}"
    try:
        ensure_number_unique(db, tenant_id, number)
    except ValueError:
        number = f"IMP-{c['nik']}-{c['row_number']}"
    ver = ed.insert_record(
        db=db, tenant_id=tenant_id, model=ContractInfo,
        identity_field="contract_id", identity_value=contract.id,
        valid_from=c["tgl_masuk"],
        valid_to=c["contract_end"] if c["contract_type"] == "PKWT" else None,
        values={"contract_type": c["contract_type"], "contract_number": number},
        event="hire", event_reason="Impor data massal", created_by=user.id,
        event_applies_to="lifecycle",
    )
    _audit("insert", "contract_info", ver.id,
           {"contract_number": number, "contract_type": c["contract_type"]})

    return person.id


def commit(db: Session, tenant_id, user, filename: str, data: bytes,
           ip: str | None = None) -> ImportCommitResponse:
    """Validasi-semua-dulu, lalu tulis atomik. Gagal satu baris = 422, DB utuh."""
    cleaned, report = validate_all(db, tenant_id, filename, data)
    if report.invalid_rows:
        raise ImportValidationError(report)
    person_ids = []
    try:
        for c in cleaned:
            person_ids.append(_write_row(db, tenant_id, user, filename, c, ip))
        db.commit()
    except Exception:
        db.rollback()
        raise
    return ImportCommitResponse(filename=filename, imported=len(person_ids),
                                person_ids=person_ids)

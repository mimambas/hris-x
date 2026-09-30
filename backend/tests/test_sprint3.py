"""Uji Sprint 3: lifecycle, validasi Indonesia, kontrak, impor Excel, dokumen."""

from __future__ import annotations

import io
import time
from datetime import date, timedelta

import openpyxl
import pytest

from app.services.imports import HEADER_ROW
from tests.conftest import login_headers

TODAY = date.today()


def ah(client):
    return login_headers(client, "hashiru", "admin_a@x.id")


_NIK_SEQ = [0]


def _nik():
    _NIK_SEQ[0] += 1
    return f"6000{_NIK_SEQ[0]:012d}"


def _person(client, h, nik=None, nama="Karyawan Uji", **kw):
    body = {"nik": nik or _nik(), "full_name": nama, "reason": "uji", **kw}
    r = client.post("/api/v1/persons", headers=h, json=body)
    assert r.status_code == 201, r.text
    return r.json()


def _employment(client, h, person_id, le_id, start="2024-01-01"):
    r = client.post(
        "/api/v1/employments",
        headers=h,
        json={
            "person_id": person_id,
            "legal_entity_id": le_id,
            "start_date": start,
            "status": "active",
            "reason": "uji",
        },
    )
    assert r.status_code == 201, r.text
    return r.json()


def _job(client, h, emp_id, ctx, event="hire", reason="Rekrutmen reguler",
         valid_from="2024-01-01"):
    return client.post(
        "/api/v1/job-info",
        headers=h,
        json={
            "employment_id": emp_id,
            "valid_from": valid_from,
            "job_id": str(ctx["job_stf"].id),
            "org_unit_id": str(ctx["ou"].id),
            "location_id": str(ctx["loc_b"].id),
            "event": event,
            "event_reason": reason,
            "reason": "uji",
        },
    )


def _comp(client, h, emp_id, event="hire", reason="Penetapan gaji awal",
          valid_from="2024-01-01", gaji=8000000):
    return client.post(
        "/api/v1/comp-info",
        headers=h,
        json={
            "employment_id": emp_id,
            "valid_from": valid_from,
            "pay_group": "Bulanan",
            "components": {"gaji_pokok": gaji, "tunjangan_tetap": 1000000},
            "event": event,
            "event_reason": reason,
            "reason": "uji",
        },
    )


def _contract(client, h, emp_id, ctype="PKWT", start="2026-01-10",
              end="2026-12-31", event="hire", reason="Rekrutmen reguler"):
    return client.post(
        "/api/v1/contracts",
        headers=h,
        json={
            "employment_id": emp_id,
            "contract_type": ctype,
            "start_date": start,
            "end_date": end,
            "event": event,
            "event_reason": reason,
            "reason": "uji",
        },
    )


# ---------------------------------------------------------------- katalog
def test_katalog_event_tersedia(client, ctx):
    h = ah(client)
    r = client.get("/api/v1/lifecycle/events", headers=h)
    assert r.status_code == 200, r.text
    codes = {e["code"] for e in r.json()}
    for code in ("hire", "promotion", "termination", "rehire", "data_update",
                 "org_founded", "org_restructure"):
        assert code in codes
    r = client.get("/api/v1/lifecycle/events",
                   params={"applies_to": "org"}, headers=h)
    assert r.status_code == 200
    assert {e["code"] for e in r.json()} <= {
        "org_founded", "org_unit_created", "org_opened", "org_renamed",
        "org_relocation", "org_restructure", "org_closed",
    }
    hire = next(e for e in client.get("/api/v1/lifecycle/events",
                                     headers=h).json() if e["code"] == "hire")
    assert "Rekrutmen reguler" in {x["reason"] for x in hire["reasons"]}


# ---------------------------------------------------------------- aturan keras event
def test_job_info_tanpa_event_valid_ditolak(client, ctx):
    h = ah(client)
    p = _person(client, h)
    e = _employment(client, h, p["id"], str(ctx["le"].id))

    r = _job(client, h, e["id"], ctx, event="", reason="")
    assert r.status_code == 422
    r = _job(client, h, e["id"], ctx, event="Ngawur", reason="Lainnya")
    assert r.status_code == 422
    r = _job(client, h, e["id"], ctx, event="hire", reason="Alasan ngawur")
    assert r.status_code == 422
    # Case-insensitive: "Hire" diterima, disimpan sebagai kode "hire".
    r = _job(client, h, e["id"], ctx, event="Hire", reason="Rekrutmen reguler")
    assert r.status_code == 201, r.text
    assert r.json()["record"]["event"] == "hire"


def test_comp_info_event_tak_valid_ditolak(client, ctx):
    h = ah(client)
    p = _person(client, h)
    e = _employment(client, h, p["id"], str(ctx["le"].id))
    r = _comp(client, h, e["id"], event="hire", reason="Ngawur")
    assert r.status_code == 422
    r = _comp(client, h, e["id"], event="bonus", reason="Lainnya")
    assert r.status_code == 422


def test_correct_event_divalidasi_katalog(client, ctx):
    h = ah(client)
    p = _person(client, h)
    e = _employment(client, h, p["id"], str(ctx["le"].id))
    rec_id = _job(client, h, e["id"], ctx).json()["record"]["id"]
    r = client.patch(
        f"/api/v1/job-info/{rec_id}/correct",
        headers=h,
        json={"event": "Ngawur", "reason": "uji"},
    )
    assert r.status_code == 422
    r = client.patch(
        f"/api/v1/job-info/{rec_id}/correct",
        headers=h,
        json={"event": "mutation", "event_reason": "Rotasi internal",
              "reason": "uji"},
    )
    assert r.status_code == 200, r.text
    assert r.json()["event"] == "mutation"


# ---------------------------------------------------------------- derivasi status
def test_terminasi_menutup_employment(client, ctx):
    h = ah(client)
    p = _person(client, h)
    e = _employment(client, h, p["id"], str(ctx["le"].id))
    assert _job(client, h, e["id"], ctx).status_code == 201
    tgl = (TODAY - timedelta(days=10)).isoformat()
    r = _job(client, h, e["id"], ctx, event="termination",
             reason="Pengunduran diri", valid_from=tgl)
    assert r.status_code == 201, r.text
    emps = client.get("/api/v1/employments", headers=h).json()
    emp = next(x for x in emps if x["id"] == e["id"])
    assert emp["status"] == "terminated"
    assert emp["end_date"] == tgl


def test_rehire_person_yang_sama(client, ctx):
    h = ah(client)
    p = _person(client, h, nama="Pegawai Rehire")
    e1 = _employment(client, h, p["id"], str(ctx["le"].id), start="2022-01-01")
    assert _job(client, h, e1["id"], ctx, valid_from="2022-01-01").status_code == 201
    tgl_out = (TODAY - timedelta(days=40)).isoformat()
    assert _job(client, h, e1["id"], ctx, event="termination",
                reason="Kontrak berakhir", valid_from=tgl_out).status_code == 201
    # Rehire = employment BARU pada person yang sama (NIK tetap).
    tgl_in = (TODAY - timedelta(days=5)).isoformat()
    e2 = _employment(client, h, p["id"], str(ctx["le"].id), start=tgl_in)
    assert e2["id"] != e1["id"]
    assert e2["person_id"] == p["id"]
    r = _job(client, h, e2["id"], ctx, event="rehire",
             reason="Rekrutmen ulang karyawan lama", valid_from=tgl_in)
    assert r.status_code == 201, r.text
    emps = {x["id"]: x for x in client.get("/api/v1/employments", headers=h).json()}
    assert emps[e1["id"]]["status"] == "terminated"
    assert emps[e2["id"]]["status"] == "active"


# ---------------------------------------------------------------- validasi Indonesia
def test_nik_duplikat_dan_pendek_ditolak(client, ctx):
    h = ah(client)
    nik = _nik()
    assert client.post("/api/v1/persons", headers=h,
                       json={"nik": nik, "full_name": "A",
                             "reason": "uji"}).status_code == 201
    r = client.post("/api/v1/persons", headers=h,
                    json={"nik": nik, "full_name": "B", "reason": "uji"})
    assert r.status_code == 409
    r = client.post("/api/v1/persons", headers=h,
                    json={"nik": "123456789012345", "full_name": "C",
                          "reason": "uji"})
    assert r.status_code == 422  # 15 digit


def test_validasi_identitas_indonesia(client, ctx):
    h = ah(client)
    base = {"nik": _nik(), "full_name": "Validasi ID", "reason": "uji"}
    # NPWP 15 digit ditolak
    r = client.post("/api/v1/persons", headers=h,
                    json={**base, "npwp": "123456789012345"})
    assert r.status_code == 422
    # PTKP tak dikenal ditolak
    r = client.post("/api/v1/persons", headers=h,
                    json={**base, "ptkp": "XX/9"})
    assert r.status_code == 422
    # Email invalid ditolak
    r = client.post("/api/v1/persons", headers=h,
                    json={**base, "email": "bukan-email"})
    assert r.status_code == 422
    # Tgl lahir masa depan ditolak
    r = client.post("/api/v1/persons", headers=h,
                    json={**base, "birth_date": (TODAY + timedelta(days=1)).isoformat()})
    assert r.status_code == 422
    # Data lengkap valid -> 201, ternormalisasi
    r = client.post("/api/v1/persons", headers=h, json={
        **base, "birth_place": "Jakarta", "birth_date": "1990-05-20",
        "email": "Valid@Contoh.ID", "npwp": "12.345.678.9012.3456",
        "ptkp": "K/1", "bpjs_kes_no": "0001234567890",
        "bpjs_tk_no": "12345678901", "bank_name": "BCA",
        "no_rekening": "1234567890",
    })
    assert r.status_code == 201, r.text
    got = r.json()
    assert got["npwp"] == "1234567890123456"
    assert got["email"] == "valid@contoh.id"
    assert got["ptkp"] == "K/1"
    # Email duplikat -> 409
    r = client.post("/api/v1/persons", headers=h,
                    json={"nik": _nik(), "full_name": "Duplikat Email",
                          "email": "valid@contoh.id", "reason": "uji"})
    assert r.status_code == 409


def test_patch_person(client, ctx):
    h = ah(client)
    p = _person(client, h, nama="Patch Me")
    r = client.patch(f"/api/v1/persons/{p['id']}", headers=h, json={
        "full_name": "Patched", "bank_name": "Mandiri", "reason": "uji"})
    assert r.status_code == 200, r.text
    assert r.json()["full_name"] == "Patched"
    assert r.json()["bank_name"] == "Mandiri"
    # PTKP tak bisa via PATCH
    r = client.patch(f"/api/v1/persons/{p['id']}", headers=h, json={
        "ptkp": "K/1", "reason": "uji"})
    assert r.status_code == 422  # field tak dikenal skema


# ---------------------------------------------------------------- PTKP
def test_ptkp_change_membuat_versi_comp_baru(client, ctx):
    h = ah(client)
    p = _person(client, h, nama="PTKP Uji", ptkp="TK/0")
    e = _employment(client, h, p["id"], str(ctx["le"].id), start="2024-01-01")
    assert _job(client, h, e["id"], ctx, valid_from="2024-01-01").status_code == 201
    assert _comp(client, h, e["id"], valid_from="2024-01-01").status_code == 201

    eff = TODAY.isoformat()
    r = client.post(f"/api/v1/persons/{p['id']}/ptkp-change", headers=h, json={
        "ptkp": "K/1", "effective_date": eff, "reason": "Menikah"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["ptkp_lama"] == "TK/0" and body["ptkp_baru"] == "K/1"
    assert len(body["comp_versions"]) == 1
    ver = body["comp_versions"][0]
    assert ver["ptkp"] == "K/1" and ver["event"] == "data_update"
    assert ver["event_reason"] == "Perubahan status PTKP"
    assert ver["components"]["gaji_pokok"] == 8000000  # nominal disalin

    # Person ikut berubah; riwayat comp: sebelum efektif masih TK/0.
    assert client.get(f"/api/v1/persons/{p['id']}", headers=h).json()["ptkp"] == "K/1"
    tl = client.get("/api/v1/comp-info/timeline",
                    params={"employment_id": e["id"]}, headers=h).json()
    assert len(tl) == 2
    kmrn = client.get("/api/v1/comp-info",
                      params={"employment_id": e["id"],
                              "as_of": (TODAY - timedelta(days=1)).isoformat()},
                      headers=h).json()
    assert kmrn["ptkp"] == "TK/0"


# ---------------------------------------------------------------- kontrak
def test_kontrak_validasi_dasar(client, ctx):
    h = ah(client)
    p = _person(client, h)
    e = _employment(client, h, p["id"], str(ctx["le"].id))
    # PKWT tanpa end_date -> 422
    r = _contract(client, h, e["id"], end=None)
    assert r.status_code == 422
    # PKWTT dengan end_date -> 422
    r = _contract(client, h, e["id"], ctype="PKWTT", end="2027-01-01")
    assert r.status_code == 422
    # Durasi > 60 bulan -> 422
    r = _contract(client, h, e["id"], start="2026-01-01", end="2032-01-01")
    assert r.status_code == 422
    # Valid
    r = _contract(client, h, e["id"])
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["current_version"]["contract_type"] == "PKWT"
    assert body["current_version"]["valid_to"] == "2026-12-31"
    assert body["current_version"]["contract_number"]
    # Nomor duplikat -> 422
    dup_no = body["current_version"]["contract_number"]
    p2 = _person(client, h)
    e2 = _employment(client, h, p2["id"], str(ctx["le"].id))
    r = client.post("/api/v1/contracts", headers=h, json={
        "employment_id": e2["id"], "contract_type": "PKWT",
        "contract_number": dup_no, "start_date": "2026-01-10",
        "end_date": "2026-12-31", "event": "hire",
        "event_reason": "Rekrutmen reguler", "reason": "uji"})
    assert r.status_code == 422


def test_kontrak_expiring_30_14_7(client, ctx):
    h = ah(client)
    ids = []
    for days in (10, 20, 40):
        p = _person(client, h)
        e = _employment(client, h, p["id"], str(ctx["le"].id))
        end = (TODAY + timedelta(days=days)).isoformat()
        r = _contract(client, h, e["id"], start=(TODAY - timedelta(days=300)).isoformat(),
                      end=end)
        assert r.status_code == 201, r.text
        ids.append(r.json()["id"])
    r = client.get("/api/v1/contracts/expiring",
                   params={"within_days": 30}, headers=h)
    assert r.status_code == 200, r.text
    got30 = {x["contract_id"] for x in r.json()}
    assert len(got30) == 2
    r = client.get("/api/v1/contracts/expiring",
                   params={"within_days": 14}, headers=h)
    assert len(r.json()) == 1
    r = client.get("/api/v1/contracts/expiring",
                   params={"within_days": 7}, headers=h)
    assert len(r.json()) == 0
    # days_remaining terisi & terurut
    r = client.get("/api/v1/contracts/expiring",
                   params={"within_days": 30}, headers=h)
    days_list = [x["days_remaining"] for x in r.json()]
    assert days_list == sorted(days_list) and set(days_list) == {10, 20}


def test_kontrak_extend_dan_convert(client, ctx):
    h = ah(client)
    p = _person(client, h)
    e = _employment(client, h, p["id"], str(ctx["le"].id), start="2024-01-01")
    assert _job(client, h, e["id"], ctx, valid_from="2024-01-01").status_code == 201
    r = _contract(client, h, e["id"], start="2024-01-01", end="2024-12-31")
    assert r.status_code == 201, r.text
    cid = r.json()["id"]

    # Extend 1x OK (batas default max_extensions=1)
    r = client.post(f"/api/v1/contracts/{cid}/extend", headers=h, json={
        "new_end_date": "2025-12-31", "event_reason": "Perpanjangan PKWT",
        "reason": "uji"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert len(body["versions"]) == 2
    assert body["current_version"]["valid_from"] == "2025-01-01"
    assert body["current_version"]["event"] == "contract_extension"
    # Sidecar di timeline job
    tl = client.get("/api/v1/job-info/timeline",
                    params={"employment_id": e["id"]}, headers=h).json()
    assert any(v["event"] == "contract_extension" for v in tl)

    # Extend ke-2 ditolak (batas kebijakan)
    r = client.post(f"/api/v1/contracts/{cid}/extend", headers=h, json={
        "new_end_date": "2026-12-31", "event_reason": "Perpanjangan PKWT",
        "reason": "uji"})
    assert r.status_code == 422

    # Convert PKWT -> PKWTT pada kontrak lain
    p2 = _person(client, h)
    e2 = _employment(client, h, p2["id"], str(ctx["le"].id), start="2024-01-01")
    assert _job(client, h, e2["id"], ctx, valid_from="2024-01-01").status_code == 201
    r = _contract(client, h, e2["id"], start="2024-01-01", end="2024-12-31")
    cid2 = r.json()["id"]
    r = client.post(f"/api/v1/contracts/{cid2}/convert", headers=h, json={
        "event_reason": "Konversi PKWT ke PKWTT", "reason": "uji"})
    assert r.status_code == 200, r.text
    cur = r.json()["current_version"]
    assert cur["contract_type"] == "PKWTT"
    assert cur["valid_to"] == "9999-12-31"
    assert cur["event"] == "contract_conversion"
    # Convert lagi -> 422 (sudah PKWTT)
    r = client.post(f"/api/v1/contracts/{cid2}/convert", headers=h, json={
        "event_reason": "Konversi PKWT ke PKWTT", "reason": "uji"})
    assert r.status_code == 422
    # PKWTT tak muncul di expiring
    exp = client.get("/api/v1/contracts/expiring",
                     params={"within_days": 365}, headers=h).json()
    assert cid2 not in {x["contract_id"] for x in exp}


def test_contract_policy_crud(client, ctx):
    h = ah(client)
    r = client.get("/api/v1/contracts/policy", headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["max_pkwt_months"] == 60
    assert r.json()["max_extensions"] == 1
    r = client.put("/api/v1/contracts/policy", headers=h, json={
        "max_pkwt_months": 24, "reason": "uji kebijakan"})
    assert r.status_code == 200, r.text
    assert r.json()["max_pkwt_months"] == 24
    # Kembalikan default agar tak memengaruhi test lain
    client.put("/api/v1/contracts/policy", headers=h, json={
        "max_pkwt_months": 60, "reason": "reset"})


# ---------------------------------------------------------------- impor Excel
def _sample_rows(n, nik_prefix="3174"):
    rows = []
    for i in range(n):
        rows.append({
            "nik": f"{nik_prefix}{i:012d}",
            "nama": f"Karyawan {i}",
            "email": f"karyawan{i}@contoh.id",
            "tgl_lahir": date(1990, 1, 1),
            "npwp": "1234567890123456",
            "ptkp": "TK/0",
            "bpjs_kes": "0001234567890",
            "bpjs_tk": "12345678901",
            "bank": "BCA",
            "no_rekening": "1234567890",
            "legal_entity": "PT Hashiru",
            "org_unit": "Eng",
            "job_code": "STF",
            "position": "Staff",
            "tgl_masuk": date(2024, 3, 1),
            "contract_type": "PKWTT",
            "contract_end": None,
            "gaji_pokok": 8000000,
            "tunjangan_tetap": 2000000,
            "lokasi": "Loc A",
        })
    return rows


def _workbook_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Karyawan"
    for ci, key in enumerate(HEADER_ROW, start=1):
        ws.cell(row=1, column=ci, value=key)
    for ri, row in enumerate(rows, start=2):
        for ci, key in enumerate(HEADER_ROW, start=1):
            val = row[key]
            cell = ws.cell(row=ri, column=ci, value=val)
            if key in ("nik", "npwp", "bpjs_kes", "bpjs_tk", "no_rekening"):
                cell.number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _person_count(client, h):
    return len(client.get("/api/v1/persons", headers=h).json())


def test_template_dapat_diunduh(client, ctx):
    h = ah(client)
    r = client.get("/api/v1/imports/employees/template", headers=h)
    assert r.status_code == 200, r.text
    assert r.content[:2] == b"PK"  # signature zip xlsx
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert "Karyawan" in wb.sheetnames and "Panduan" in wb.sheetnames
    headers = [c.value for c in next(wb["Karyawan"].iter_rows(min_row=1, max_row=1))]
    assert headers == HEADER_ROW


def test_dry_run_500_baris_dengan_3_error(client, ctx):
    h = ah(client)
    rows = _sample_rows(500)
    data = _workbook_bytes(rows)
    # Suntik 3 error: NIK pendek (baris 10), job tak dikenal (200), NIK duplikat (400).
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    col = {k: i + 1 for i, k in enumerate(HEADER_ROW)}
    ws.cell(row=10, column=col["nik"], value="12345")
    ws.cell(row=200, column=col["job_code"], value="XXX")
    ws.cell(row=400, column=col["nik"], value=rows[4]["nik"])  # duplikat data-ke-5
    buf = io.BytesIO()
    wb.save(buf)
    bad = buf.getvalue()

    before = _person_count(client, h)
    t0 = time.perf_counter()
    r = client.post(
        "/api/v1/imports/employees/dry-run",
        headers=h,
        files={"file": ("karyawan.xlsx", bad,
                        "application/vnd.openxmlformats-officedocument"
                        ".spreadsheetml.sheet")},
    )
    dt = time.perf_counter() - t0
    assert r.status_code == 200, r.text
    report = r.json()
    assert report["total_rows"] == 500
    assert report["invalid_rows"] == 3, report["rows"]
    assert report["valid_rows"] == 497
    by_row = {x["row_number"]: x for x in report["rows"]}
    assert set(by_row) == {10, 200, 400}
    assert by_row[10]["errors"][0]["field"] == "nik"
    assert by_row[200]["errors"][0]["field"] == "job_code"
    assert by_row[400]["errors"][0]["field"] == "nik"
    # Dry-run TIDAK menulis DB
    assert _person_count(client, h) == before
    print(f"\ndry-run 500 baris: {dt:.2f} dtk")


def test_commit_atomic_rollback_dan_sukses(client, ctx):
    h = ah(client)
    rows = _sample_rows(3, nik_prefix="3175")
    bad = _workbook_bytes(rows)
    wb = openpyxl.load_workbook(io.BytesIO(bad))
    ws = wb.active
    col = {k: i + 1 for i, k in enumerate(HEADER_ROW)}
    ws.cell(row=3, column=col["nik"], value="pendek")
    buf = io.BytesIO()
    wb.save(buf)
    bad = buf.getvalue()

    before = _person_count(client, h)
    files = {"file": ("karyawan.xlsx", bad,
                      "application/vnd.openxmlformats-officedocument"
                      ".spreadsheetml.sheet")}
    r = client.post("/api/v1/imports/employees/commit", headers=h, files=files)
    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert detail["invalid_rows"] == 1
    assert detail["rows"][0]["row_number"] == 3
    # Atomic: tak ada yang tertulis
    assert _person_count(client, h) == before

    # Perbaiki -> commit sukses
    good = _workbook_bytes(rows)
    r = client.post(
        "/api/v1/imports/employees/commit", headers=h,
        files={"file": ("karyawan.xlsx", good,
                        "application/vnd.openxmlformats-officedocument"
                        ".spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    assert r.json()["imported"] == 3
    assert _person_count(client, h) == before + 3
    # Data terimpor benar: cek salah satu
    persons = client.get("/api/v1/persons", headers=h).json()
    imported = next(p for p in persons if p["nik"] == rows[0]["nik"])
    assert imported["full_name"] == "Karyawan 0"
    assert imported["npwp"] == "1234567890123456"
    # Audit channel=import tercatat
    logs = client.get("/api/v1/audit-logs",
                      params={"object_type": "person",
                              "object_id": imported["id"]}, headers=h).json()
    assert logs and all(l["channel"] == "import" for l in logs)


def test_commit_500_baris_sukses(client, ctx):
    h = ah(client)
    data = _workbook_bytes(_sample_rows(500, nik_prefix="3176"))
    before = _person_count(client, h)
    t0 = time.perf_counter()
    r = client.post(
        "/api/v1/imports/employees/commit", headers=h,
        files={"file": ("karyawan_500.xlsx", data,
                        "application/vnd.openxmlformats-officedocument"
                        ".spreadsheetml.sheet")},
    )
    dt = time.perf_counter() - t0
    assert r.status_code == 200, r.text
    assert r.json()["imported"] == 500
    assert _person_count(client, h) == before + 500
    print(f"\ncommit 500 baris: {dt:.2f} dtk")


# ---------------------------------------------------------------- dokumen
def test_dokumen_upload_unduh_versi(client, ctx):
    h = ah(client)
    p = _person(client, h, nama="Dokumen Uji")
    pdf = b"%PDF-1.4-contoh"

    def _upload(content, fname="ktp.pdf"):
        return client.post(
            "/api/v1/documents", headers=h,
            data={"doc_type": "ktp", "person_id": p["id"],
                  "notes": "KTP uji"},
            files={"file": (fname, content, "application/pdf")},
        )

    r = _upload(pdf)
    assert r.status_code == 201, r.text
    d1 = r.json()
    assert d1["version"] == 1 and d1["is_current"] is True

    r = _upload(b"%PDF-1.4-revisi")
    assert r.status_code == 201, r.text
    d2 = r.json()
    assert d2["version"] == 2 and d2["is_current"] is True

    # Default: hanya versi kini
    docs = client.get("/api/v1/documents",
                      params={"person_id": p["id"]}, headers=h).json()
    assert len(docs) == 1 and docs[0]["version"] == 2
    docs = client.get("/api/v1/documents",
                      params={"person_id": p["id"], "include_old": True},
                      headers=h).json()
    assert len(docs) == 2

    # Unduh versi terbaru
    r = client.get(f"/api/v1/documents/{d2['id']}/download", headers=h)
    assert r.status_code == 200
    assert r.content == b"%PDF-1.4-revisi"

    # Ekstensi terlarang ditolak
    r = client.post(
        "/api/v1/documents", headers=h,
        data={"doc_type": "ktp", "person_id": p["id"]},
        files={"file": ("jahat.exe", b"MZ", "application/octet-stream")},
    )
    assert r.status_code == 422
    # Tanpa subjek ditolak
    r = client.post(
        "/api/v1/documents", headers=h,
        data={"doc_type": "ktp"},
        files={"file": ("ktp.pdf", pdf, "application/pdf")},
    )
    assert r.status_code == 422


def test_dokumen_rbp_ditolak_tanpa_izin(client, ctx):
    # u_none: role Empty tanpa izin apa pun -> 403
    h_none = login_headers(client, "hashiru", "u_none@x.id")
    r = client.get("/api/v1/documents", headers=h_none)
    assert r.status_code == 403

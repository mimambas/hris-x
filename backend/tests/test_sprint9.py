"""Sprint 9 (PRD 24.2 S9, ANL-001 s/d ANL-004): dasbor & laporan standar.

Demo goal PRD: "Direktur melihat headcount dan turnover real-time."
"""

from __future__ import annotations

from datetime import date, datetime, timezone

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import (
    AttendanceRecord,
    Employment,
    FieldPermission,
    LeaveBalance,
    LeaveRequest,
    LeaveType,
    PayrollLine,
    PayrollRun,
    PermissionRole,
)

from .conftest import PASSWORD, login_headers


def ah(client):
    return login_headers(client, "hashiru", "admin_a@x.id")


def mh(client):
    return login_headers(client, "hashiru", "u_mgr@x.id")


def nh(client):
    return login_headers(client, "hashiru", "u_none@x.id")


_NIK_SEQ = [900]


def _nik():
    _NIK_SEQ[0] += 1
    return f"9100{_NIK_SEQ[0]:012d}"


def _mk9(client, h, ctx, nama, gender=None, start="2024-01-01",
         manager_emp_id=None, birth_date=None):
    """Person + employment + job-info (di unit 'Eng' fixture)."""
    payload = {"nik": _nik(), "full_name": nama, "reason": "uji sprint9"}
    if gender:
        payload["gender"] = gender
    if birth_date:
        payload["birth_date"] = birth_date
    p = client.post("/api/v1/persons", headers=h, json=payload).json()
    assert "id" in p, p
    e = client.post("/api/v1/employments", headers=h, json={
        "person_id": p["id"], "legal_entity_id": str(ctx["le"].id),
        "start_date": start, "status": "active",
        "reason": "uji sprint9"}).json()
    job_payload = {"employment_id": e["id"], "valid_from": start,
                   "job_id": str(ctx["job_stf"].id),
                   "org_unit_id": str(ctx["ou"].id),
                   "location_id": str(ctx["loc_b"].id),
                   "event": "hire", "event_reason": "Rekrutmen reguler",
                   "reason": "uji sprint9"}
    if manager_emp_id:
        job_payload["manager_employment_id"] = str(manager_emp_id)
    r = client.post("/api/v1/job-info", headers=h, json=job_payload)
    assert r.status_code == 201, r.text
    return p, e


def _grant_dashboard(ctx):
    """Izin dashboard view untuk role MgrRole (fixture)."""
    db, ta = ctx["db"], ctx["ta"]
    r_mgr = db.execute(select(PermissionRole).where(
        PermissionRole.tenant_id == ta.id,
        PermissionRole.name == "MgrRole")).scalar_one()
    db.add(FieldPermission(tenant_id=ta.id, role_id=r_mgr.id,
                           object_name="dashboard", field_name="*",
                           can_view=True))
    db.commit()


def _terminate(client, h, ctx, emp_id, end):
    r = client.post("/api/v1/job-info", headers=h, json={
        "employment_id": emp_id, "valid_from": end,
        "job_id": str(ctx["job_stf"].id), "org_unit_id": str(ctx["ou"].id),
        "location_id": str(ctx["loc_b"].id),
        "event": "termination", "event_reason": "Pengunduran diri",
        "reason": "uji sprint9"})
    assert r.status_code == 201, r.text


# ---------------------------------------------------------------- gender
def test_gender_dinormalisasi_dan_divalidasi(client, ctx):
    h = ah(client)
    r = client.post("/api/v1/persons", headers=h, json={
        "nik": _nik(), "full_name": "Gender Uji", "gender": "l",
        "reason": "uji"})
    assert r.status_code == 201, r.text
    assert r.json()["gender"] == "L"
    pid = r.json()["id"]
    r = client.patch(f"/api/v1/persons/{pid}", headers=h, json={
        "gender": "X", "reason": "uji"})
    assert r.status_code == 422, r.text


# ---------------------------------------------------------------- headcount
def test_headcount_total_dan_breakdown(client, ctx):
    h = ah(client)
    _mk9(client, h, ctx, "HC Satu", gender="L")
    _mk9(client, h, ctx, "HC Dua", gender="P")
    # 4 karyawan fixture (tanpa gender) + 2 baru = 6
    r = client.get("/api/v1/dashboard/headcount?as_of=2024-06-01", headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["total"] == 6
    assert d["by_gender"] == {"Tidak diisi": 4, "L": 1, "P": 1}
    assert d["by_org_unit"] == {"Eng": 6}
    assert d["by_contract_type"] == {"Tanpa kontrak": 6}
    assert d["by_status"] == {"active": 6}
    assert d["new_this_month"] == 0
    assert d["left_this_month"] == 0


def test_headcount_effective_dating_dan_kontrak(client, ctx):
    h = ah(client)
    p, e = _mk9(client, h, ctx, "HC Kontrak", gender="L", start="2024-02-15")
    r = client.post("/api/v1/contracts", headers=h, json={
        "employment_id": e["id"], "contract_type": "PKWT",
        "contract_number": "PKWT/001", "start_date": "2024-02-15",
        "end_date": "2025-02-14", "event": "hire",
        "event_reason": "Rekrutmen reguler", "reason": "uji"})
    assert r.status_code == 201, r.text
    # Sebelum 2024-02-15: hanya 4 fixture (start 2024-01-01).
    r = client.get("/api/v1/dashboard/headcount?as_of=2024-02-01", headers=h)
    assert r.json()["total"] == 4
    r = client.get("/api/v1/dashboard/headcount?as_of=2024-02-15", headers=h)
    d = r.json()
    assert d["total"] == 5
    assert d["by_contract_type"] == {"PKWT": 1, "Tanpa kontrak": 4}
    assert d["new_this_month"] == 1


def test_headcount_karyawan_biasa_403(client, ctx):
    r = client.get("/api/v1/dashboard/headcount", headers=nh(client))
    assert r.status_code == 403


# ---------------------------------------------------------------- turnover
def test_turnover_rate_sesuai_anl004(client, ctx):
    h = ah(client)
    _mk9(client, h, ctx, "TO Tetap", gender="L")
    _, e = _mk9(client, h, ctx, "TO Keluar", gender="P")
    _terminate(client, h, ctx, e["id"], "2024-03-20")
    r = client.get("/api/v1/dashboard/turnover?period=2024-03", headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["terminated"] == 1
    assert d["headcount_start"] == 6   # 4 fixture + 2 baru
    assert d["headcount_end"] == 5     # 1 keluar
    assert d["avg_headcount"] == 5.5
    assert d["rate_pct"] == round(1 / 5.5 * 100, 2)
    assert d["by_org_unit"] == {"Eng": 1}
    assert len(d["trend_12_months"]) == 12
    assert d["trend_12_months"][-1]["period"] == "2024-03"
    assert d["trend_12_months"][-1]["terminated"] == 1


def test_turnover_nol_bila_tanpa_terminasi(client, ctx):
    h = ah(client)
    r = client.get("/api/v1/dashboard/turnover?period=2024-04", headers=h)
    d = r.json()
    assert d["terminated"] == 0
    assert d["rate_pct"] == 0.0


def test_format_period_salah_422(client, ctx):
    h = ah(client)
    for path in ("/api/v1/dashboard/turnover?period=2024",
                 "/api/v1/dashboard/attendance?period=kemarin",
                 "/api/v1/reports/payroll-summary.xlsx?period=2024-13"):
        r = client.get(path, headers=h)
        assert r.status_code == 422, (path, r.text)


# ---------------------------------------------------------------- attendance
def _att(db, ta, emp_id, day, check_in, check_out, late, work_min,
         status="present"):
    db.add(AttendanceRecord(
        tenant_id=ta.id, employment_id=emp_id, date=date(2024, 5, day),
        source="web",
        check_in=(datetime(2024, 5, day, *check_in, tzinfo=timezone.utc)
                  if check_in else None),
        check_out=(datetime(2024, 5, day, *check_out, tzinfo=timezone.utc)
                   if check_out else None),
        status=status, late_minutes=late, work_minutes=work_min))


def test_dashboard_absensi(client, ctx):
    h = ah(client)
    db, ta = ctx["db"], ctx["ta"]
    _, e1 = _mk9(client, h, ctx, "AB Satu", gender="L")
    _, e2 = _mk9(client, h, ctx, "AB Dua", gender="P")
    import uuid
    eid1, eid2 = uuid.UUID(e1["id"]), uuid.UUID(e2["id"])
    _att(db, ta, eid1, 6, (8, 0), (17, 0), 0, 480)     # tepat waktu
    _att(db, ta, eid1, 7, (8, 45), (17, 0), 45, 435)    # telat 45 mnt
    _att(db, ta, eid2, 6, (8, 0), (17, 0), 0, 480)
    _att(db, ta, eid2, 7, None, None, 0, 0, status="absent")
    db.commit()
    r = client.get("/api/v1/dashboard/attendance?period=2024-05", headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["check_ins"] == 3
    assert d["late_count"] == 1
    assert d["late_pct"] == round(1 / 3 * 100, 2)
    assert d["absent_days"] == 1
    assert d["avg_work_hours"] == round((480 + 435 + 480) / 60 / 2, 2)
    assert d["by_org_unit"]["Eng"]["late"] == 1


# ---------------------------------------------------------------- leave
def test_dashboard_cuti(client, ctx):
    h = ah(client)
    db, ta = ctx["db"], ctx["ta"]
    _, e = _mk9(client, h, ctx, "CT Satu", gender="L")
    import uuid
    eid = uuid.UUID(e["id"])
    lt = LeaveType(tenant_id=ta.id, name="Cuti Tahunan", code="CT",
                   quota_days=12)
    db.add(lt)
    db.flush()
    db.add(LeaveRequest(tenant_id=ta.id, employment_id=eid,
                        leave_type_id=lt.id, start_date=date(2024, 7, 1),
                        end_date=date(2024, 7, 3), days=3, status="approved"))
    db.add(LeaveBalance(tenant_id=ta.id, employment_id=eid,
                        leave_type_id=lt.id, year=2024, entitled=12,
                        used=3, remaining=9))
    db.commit()
    r = client.get("/api/v1/dashboard/leave?year=2024", headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["approved_requests"] == 1
    assert d["total_leave_days"] == 3
    assert d["top_units_by_leave_days"][0] == {"org_unit": "Eng", "days": 3}
    bal = next(b for b in d["balances_by_type"]
               if b["leave_type"] == "Cuti Tahunan")
    assert (bal["entitled"], bal["used"], bal["remaining"]) == (12, 3, 9)


# ---------------------------------------------------------------- payroll
def _payroll_run(db, ta, ctx, period="2024-08", lines_spec=()):
    run = PayrollRun(tenant_id=ta.id, period=period, status="locked")
    db.add(run)
    db.flush()
    import uuid
    for spec in lines_spec:
        db.add(PayrollLine(
            tenant_id=ta.id, payroll_run_id=run.id,
            employment_id=uuid.UUID(spec["emp_id"]),
            person_name=spec["nama"], nik=spec["nik"],
            gross=spec["gross"], total_deductions=spec["ded"],
            pph21=spec["pph21"],
            reimbursement_amount=spec.get("reimb", 0),
            take_home_pay=spec["gross"] - spec["ded"] - spec["pph21"]
            + spec.get("reimb", 0),
            breakdown={"lembur": spec.get("lembur", 0)},
            bank_name="BCA", bank_account_no="9000000001"))
    db.commit()
    return run


def test_dashboard_payroll_dan_xlsx(client, ctx):
    h = ah(client)
    db, ta = ctx["db"], ctx["ta"]
    p1, e1 = _mk9(client, h, ctx, "PR Satu", gender="L")
    p2, e2 = _mk9(client, h, ctx, "PR Dua", gender="P")
    _payroll_run(db, ta, ctx, lines_spec=[
        {"emp_id": e1["id"], "nama": "PR Satu", "nik": p1["nik"],
         "gross": 10_000_000, "ded": 1_000_000, "pph21": 500_000,
         "lembur": 200_000, "reimb": 100_000},
        {"emp_id": e2["id"], "nama": "PR Dua", "nik": p2["nik"],
         "gross": 8_000_000, "ded": 800_000, "pph21": 300_000},
    ])
    r = client.get("/api/v1/dashboard/payroll?period=2024-08", headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    t = d["totals"]
    assert t["gross"] == 18_000_000
    assert t["deductions"] == 1_800_000
    assert t["pph21"] == 800_000
    assert t["reimbursement"] == 100_000
    assert t["overtime"] == 200_000
    assert t["take_home"] == 15_500_000
    assert d["by_org_unit"]["Eng"]["gross"] == 18_000_000
    # 404 untuk periode tanpa run.
    r = client.get("/api/v1/dashboard/payroll?period=2024-09", headers=h)
    assert r.status_code == 404
    # XLSX ringkasan payroll.
    r = client.get("/api/v1/reports/payroll-summary.xlsx?period=2024-08",
                   headers=h)
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(filename=__import__("io").BytesIO(r.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == [
        "NIK", "Nama", "Unit", "Bruto", "Potongan", "PPh 21",
        "Reimbursement", "Take-Home", "Bank", "Rekening"]
    assert ws.max_row == 3  # header + 2 baris


def test_dashboard_payroll_karyawan_biasa_403(client, ctx):
    r = client.get("/api/v1/dashboard/payroll?period=2024-08",
                   headers=nh(client))
    assert r.status_code == 403


# ---------------------------------------------------------------- demographics
def test_dashboard_demografi(client, ctx):
    h = ah(client)
    _mk9(client, h, ctx, "DG Muda", gender="L", birth_date="2000-05-01")
    _mk9(client, h, ctx, "DG Senior", gender="P", birth_date="1975-03-10")
    r = client.get("/api/v1/dashboard/demographics?as_of=2024-06-01",
                   headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["total"] == 6
    assert d["by_age"]["20-29"] == 1
    assert d["by_age"]["40-49"] == 1
    assert d["by_gender"] == {"Tidak diisi": 4, "L": 1, "P": 1}
    # Semua mulai 2024-01-01 -> masa kerja <1 tahun per 2024-06-01.
    assert d["by_tenure"] == {"<1 tahun": 6}


# ---------------------------------------------------------------- RBP / population
def test_manajer_hanya_melihat_timnya(client, ctx):
    _grant_dashboard(ctx)
    h = ah(client)
    mh_ = mh(client)
    mgr_emp_id = str(ctx["e_mgr"].id)
    _mk9(client, h, ctx, "Tim A", gender="L", manager_emp_id=mgr_emp_id)
    _mk9(client, h, ctx, "Bukan Tim", gender="P")
    r = client.get("/api/v1/dashboard/headcount", headers=mh_)
    assert r.status_code == 200, r.text
    assert r.json()["total"] == 1
    r = client.get("/api/v1/dashboard/headcount", headers=h)
    assert r.json()["total"] >= 3  # admin: semua (4 fixture + 2 baru)


def test_reports_employees_xlsx_dan_populasi(client, ctx):
    _grant_dashboard(ctx)
    h = ah(client)
    mh_ = mh(client)
    mgr_emp_id = str(ctx["e_mgr"].id)
    _mk9(client, h, ctx, "XLSX Tim", gender="L", manager_emp_id=mgr_emp_id)
    r = client.get("/api/v1/reports/employees.xlsx", headers=h)
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(filename=__import__("io").BytesIO(r.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == [
        "NIK", "Nama", "Unit", "Jabatan", "Tgl Masuk", "Status",
        "Kontrak", "Jenis Kelamin"]
    rows_admin = ws.max_row - 1
    # Manajer hanya mendapat baris timnya (1 orang).
    r = client.get("/api/v1/reports/employees.xlsx", headers=mh_)
    assert r.status_code == 200, r.text
    wb2 = load_workbook(filename=__import__("io").BytesIO(r.content))
    assert wb2.active.max_row - 1 == 1
    assert rows_admin > 1
    # Karyawan biasa 403.
    r = client.get("/api/v1/reports/employees.xlsx", headers=nh(client))
    assert r.status_code == 403


def test_audit_mencatat_export(client, ctx):
    from app.models import AuditLog
    h = ah(client)
    db, ta = ctx["db"], ctx["ta"]
    before = db.execute(select(AuditLog).where(
        AuditLog.tenant_id == ta.id,
        AuditLog.action == "export")).scalars().all()
    client.get("/api/v1/reports/employees.xlsx", headers=h)
    after = db.execute(select(AuditLog).where(
        AuditLog.tenant_id == ta.id,
        AuditLog.action == "export")).scalars().all()
    assert len(after) == len(before) + 1
